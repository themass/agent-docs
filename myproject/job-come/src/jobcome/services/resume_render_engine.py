"""HTML / PDF / DOCX rendering for resume export."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document
from jinja2 import Environment, FileSystemLoader, select_autoescape

_TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "templates" / "resume"


class ResumeRenderEngine:
    def __init__(self) -> None:
        self._env = Environment(
            loader=FileSystemLoader(str(_TEMPLATE_DIR)),
            autoescape=select_autoescape(["html", "xml"]),
        )

    def render_html(self, sections: dict[str, Any], *, template_id: str = "zh-one-page") -> str:
        template = self._env.get_template(f"{template_id}.html")
        if template_id == "zh-en-bilingual":
            return template.render(
                zh=sections.get("zh_sections") or {},
                en=sections.get("en_sections") or {},
            )
        return template.render(**sections)

    def render_pdf(self, html: str) -> bytes:
        try:
            from playwright.sync_api import sync_playwright
        except ImportError as exc:
            raise RuntimeError("playwright not installed; pip install 'jobcome[export]'") from exc

        with sync_playwright() as p:
            browser = p.chromium.launch()
            try:
                page = browser.new_page()
                page.set_content(html, wait_until="networkidle")
                return page.pdf(format="A4", print_background=True)
            finally:
                browser.close()

    def render_docx(self, sections: dict[str, Any]) -> bytes:
        if sections.get("track") == "bilingual":
            buffer = io.BytesIO()
            doc = Document()
            self._append_docx_locale(doc, sections.get("zh_sections") or {}, english=False)
            doc.add_page_break()
            self._append_docx_locale(doc, sections.get("en_sections") or {}, english=True)
            doc.save(buffer)
            return buffer.getvalue()
        doc = Document()
        self._append_docx_locale(doc, sections, english=False)
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _append_docx_locale(doc: Document, sections: dict[str, Any], *, english: bool) -> None:
        header = sections.get("header") or {}
        doc.add_heading(str(header.get("name", "")), level=0)
        if header.get("headline"):
            doc.add_paragraph(str(header["headline"]))
        if header.get("contact_line"):
            doc.add_paragraph(str(header["contact_line"]))
        summary_title = "Summary" if english else "个人总结"
        exp_title = "Experience" if english else "工作经历"
        if sections.get("summary"):
            doc.add_heading(summary_title, level=1)
            doc.add_paragraph(str(sections["summary"]))
        if sections.get("experience_blocks"):
            doc.add_heading(exp_title, level=1)
        for exp in sections.get("experience_blocks") or []:
            doc.add_heading(f"{exp.get('company')} — {exp.get('title')}", level=2)
            doc.add_paragraph(str(exp.get("date_range", "")))
            for bullet in exp.get("bullets", []):
                doc.add_paragraph(str(bullet), style="List Bullet")
