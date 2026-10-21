#!/usr/bin/env python3
"""Offline, deidentified extraction benchmark; never calls LLM or URL converters.
Run with shared Python; output is a JSON report, generated input files stay temporary.
"""

from __future__ import annotations

import argparse
import importlib.metadata
import importlib.util
import json
import tempfile
from pathlib import Path

from lib.bootstrap import ensure_project_env

PROJECT_ROOT = ensure_project_env()
SERVER_ROOT = PROJECT_ROOT / "backend" / "server"

import fitz  # noqa: E402
from docx import Document  # noqa: E402
from jc.services.document_extractor import DocumentExtractor  # noqa: E402


def samples(directory):
    root = SERVER_ROOT
    cases = []
    for file in sorted((root / "tests/regression_assets/resumes").glob("*.txt")):
        text = file.read_text()
        path = directory / file.name
        path.write_text(text)
        anchors = [line.strip() for line in text.splitlines() if line.strip()][:5]
        cases.append((path, anchors, False))
    anchors = [
        "Candidate Example",
        "Product Manager",
        "Internet Company",
        "SQL analytics",
        "University Example",
        "Revenue +30%",
    ]
    doc = Document()
    doc.add_heading(anchors[0], 0)
    table = doc.add_table(rows=3, cols=2)
    for cell, text in zip([c for row in table.rows for c in row.cells], anchors):
        cell.text = text
    path = directory / "table.docx"
    doc.save(path)
    cases.append((path, anchors, False))
    with fitz.open() as pdf:
        page = pdf.new_page()
        for i, text in enumerate(anchors):
            page.insert_text((40 if i < 3 else 320, 60 + (i % 3) * 30), text)
        path = directory / "two-column.pdf"
        path.write_bytes(pdf.tobytes())
        cases.append((path, anchors, False))
        raster = page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes("png")
    with fitz.open() as scan:
        page = scan.new_page()
        page.insert_image(page.rect, stream=raster)
        path = directory / "scanned.pdf"
        path.write_bytes(scan.tobytes())
        cases.append((path, anchors, True))
    # Chinese PDF text layer is a separate measured case.
    chinese = ["张明", "产品经理", "互联网公司", "数据分析", "本科教育"]
    with fitz.open() as pdf:
        page = pdf.new_page()
        for i, text in enumerate(chinese):
            page.insert_text((40, 50 + i * 25), text, fontname="china-s")
        path = directory / "chinese.pdf"
        path.write_bytes(pdf.tobytes())
        cases.append((path, chinese, False))
    return cases


def evaluate(output):
    from markitdown import MarkItDown

    md = MarkItDown(enable_plugins=False)
    native = DocumentExtractor()
    rows = []
    with tempfile.TemporaryDirectory(prefix="jc-docs-") as directory:
        for path, anchors, scan in samples(Path(directory)):
            for engine in ["native", "markitdown"]:
                try:
                    text = (
                        native.extract(
                            file_name=path.name, content_type=None, raw_bytes=path.read_bytes()
                        )
                        if engine == "native"
                        else md.convert(str(path)).text_content
                    )

                    def squash(value: str) -> str:
                        return "".join(value.split())

                    hits = [squash(anchor) in squash(text) for anchor in anchors]
                    positions = [squash(text).find(squash(a)) for a in anchors]
                    rows.append(
                        {
                            "file": path.name,
                            "engine": engine,
                            "scanned": scan,
                            "status": "ok",
                            "anchor_recall": sum(hits) / len(hits),
                            "anchor_order_preserved": all(
                                a <= b for a, b in zip(positions, positions[1:])
                            )
                            if all(hits)
                            else None,
                            "ocr_required": scan and not any(hits),
                            "characters": len(text),
                        }
                    )
                except Exception as exc:
                    rows.append(
                        {
                            "file": path.name,
                            "engine": engine,
                            "scanned": scan,
                            "status": "error",
                            "error": type(exc).__name__ + ": " + str(exc)[:300],
                        }
                    )
    report = {
        "scope": "offline extraction only; not structured LLM or production OCR accuracy",
        "versions": {
            "markitdown": importlib.metadata.version("markitdown"),
            "pymupdf": importlib.metadata.version("pymupdf"),
        },
        "docling": {
            "installed": bool(importlib.util.find_spec("docling")),
            "status": "not_evaluated",
            "reason": (
                "Not installed in shared environment; "
                "model/OCR downloads are not part of offline benchmark."
            ),
        },
        "cases": rows,
    }
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    ordinary = [row for row in rows if row["engine"] == "native" and not row["scanned"]]
    if any(row["status"] != "ok" or row["anchor_recall"] != 1 for row in ordinary):
        raise SystemExit("FAIL: native non-scanned extraction lost anchors")
    print(
        f"PASS: {len(rows)} extraction observations; scan OCR and Docling explicitly not accepted"
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output", type=Path, default=Path("docs/evidence/document-evaluation.json")
    )
    evaluate(parser.parse_args().output)
