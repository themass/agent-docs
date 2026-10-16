"""Regression coverage for the controlled local MarkItDown Route B."""

import io
import sys
import time
from types import SimpleNamespace

import fitz
import pytest
from docx import Document

from jc.services.document_extractor import DocumentExtractor


def test_markitdown_route_b_uses_uploaded_bytes_and_disables_plugins(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: dict[str, object] = {}

    class FakeMarkItDown:
        def __init__(self, **kwargs: object) -> None:
            calls["init"] = kwargs

        def convert_stream(self, stream, **kwargs: object):
            calls["bytes"] = stream.read()
            calls["convert"] = kwargs
            return SimpleNamespace(text_content="# Candidate\n\nExperience")

    monkeypatch.setitem(sys.modules, "markitdown", SimpleNamespace(MarkItDown=FakeMarkItDown))
    result = DocumentExtractor().extract_markdown_route_b(
        file_name="resume.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        raw_bytes=b"docx-content",
    )

    assert result == "# Candidate\n\nExperience"
    assert calls["init"] == {"enable_plugins": False}
    assert calls["bytes"] == b"docx-content"
    assert calls["convert"] == {"file_extension": ".docx"}


def test_markitdown_route_b_rejects_over_limit_input() -> None:
    extractor = DocumentExtractor()
    result = extractor.extract_markdown_route_b(
        file_name="resume.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        raw_bytes=b"x" * (extractor.MARKITDOWN_MAX_INPUT_BYTES + 1),
    )
    assert result == ""


@pytest.mark.asyncio
async def test_markitdown_route_b_timeout_does_not_block_request(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    extractor = DocumentExtractor()
    extractor.MARKITDOWN_TIMEOUT_SEC = 0.01

    def slow_convert(**kwargs: object) -> str:
        time.sleep(0.1)
        return "too late"

    monkeypatch.setattr(extractor, "extract_markdown_route_b", slow_convert)
    started = time.monotonic()
    result = await extractor.extract_markdown_route_b_with_timeout(
        file_name="resume.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        raw_bytes=b"content",
    )
    assert result == ""
    assert time.monotonic() - started < 0.08


def test_markitdown_extracts_word_table() -> None:
    document = Document()
    document.add_heading("Candidate", 0)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Company"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "JC"
    table.cell(1, 1).text = "Backend Engineer"
    stream = io.BytesIO()
    document.save(stream)

    result = DocumentExtractor().extract_markdown_route_b(
        file_name="resume-table.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        raw_bytes=stream.getvalue(),
    )
    assert "Candidate" in result
    assert "Backend Engineer" in result


def test_markitdown_extracts_text_pdf_and_degrades_on_scanned_pdf() -> None:
    text_pdf = fitz.open()
    text_page = text_pdf.new_page()
    text_page.insert_text((72, 72), "Candidate: Zhang San\nPlatform Engineer")
    text_bytes = text_pdf.tobytes()
    text_pdf.close()

    extractor = DocumentExtractor()
    text_result = extractor.extract_markdown_route_b(
        file_name="resume.pdf", content_type="application/pdf", raw_bytes=text_bytes
    )
    assert "Platform Engineer" in text_result

    scanned_pdf = fitz.open()
    page = scanned_pdf.new_page()
    image = fitz.Pixmap(fitz.csRGB, 20, 20, b"\xff" * 1_200, False)
    page.insert_image(page.rect, pixmap=image)
    scanned_bytes = scanned_pdf.tobytes()
    scanned_pdf.close()

    scanned_result = extractor.extract_markdown_route_b(
        file_name="scanned.pdf", content_type="application/pdf", raw_bytes=scanned_bytes
    )
    assert isinstance(scanned_result, str)
