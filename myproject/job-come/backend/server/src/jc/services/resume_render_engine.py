"""HTML / PDF / DOCX rendering for resume export."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from jinja2 import Environment, FileSystemLoader, select_autoescape

_TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "templates" / "resume"

_DOCX_LABELS_ZH = {
    "summary": "个人总结",
    "experience": "工作经历",
    "projects": "项目经历",
    "education": "教育背景",
    "skills": "技能",
}
_DOCX_LABELS_EN = {
    "summary": "Summary",
    "experience": "Experience",
    "projects": "Projects",
    "education": "Education",
    "skills": "Skills",
}


class ResumeRenderEngine:
    def __init__(self) -> None:
        self._env = Environment(
            loader=FileSystemLoader(str(_TEMPLATE_DIR)),
            autoescape=select_autoescape(["html", "xml"]),
        )

    def render_html(self, sections: dict[str, Any], *, template_id: str = "zh-one-page") -> str:
        template = self._env.get_template(f"{template_id}.html")
        if template_id == "zh-en-bilingual":
            return template.render(
                zh=sections.get("zh_sections") or {},
                en=sections.get("en_sections") or {},
            )
        return template.render(**sections)

    def render_pdf(self, html: str) -> bytes:
        try:
            from playwright.sync_api import sync_playwright
        except ImportError as exc:
            raise RuntimeError("playwright not installed; pip install 'jc[export]'") from exc

        with sync_playwright() as p:
            browser = p.chromium.launch()
            try:
                page = browser.new_page()
                page.route("**/*", lambda route: route.abort())
                page.set_content(html, wait_until="load")
                return page.pdf(format="A4", print_background=True)
            finally:
                browser.close()

    def render_docx(self, sections: dict[str, Any]) -> bytes:
        if sections.get("track") == "bilingual":
            buffer = io.BytesIO()
            doc = self._new_docx()
            self._append_docx_locale(doc, sections.get("zh_sections") or {}, english=False)
            doc.add_page_break()
            self._append_docx_locale(doc, sections.get("en_sections") or {}, english=True)
            doc.save(buffer)
            return buffer.getvalue()
        doc = self._new_docx()
        self._append_docx_locale(doc, sections, english=sections.get("locale", "").startswith("en"))
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _new_docx():
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(1.6)
        section.left_margin = section.right_margin = Cm(1.8)
        for name, size in [
            ("Normal", 10),
            ("Title", 22),
            ("Heading 1", 12),
            ("Heading 2", 11),
            ("List Bullet", 10),
        ]:
            style = doc.styles[name]
            style.font.name = "Arial"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            # Explicit CJK family avoids theme-font resolution losing Chinese glyphs.
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Heiti SC")
            style.paragraph_format.space_after = Pt(5)
            style.paragraph_format.space_before = Pt(8 if "Heading" in name else 0)
        return doc

    @staticmethod
    def _append_docx_locale(doc: Document, sections: dict[str, Any], *, english: bool) -> None:
        # Section order mirrors the HTML templates so PDF, HTML and Word stay consistent.
        labels = _DOCX_LABELS_EN if english else _DOCX_LABELS_ZH
        header = sections.get("header") or {}
        doc.add_heading(str(header.get("name", "")), level=0)
        if header.get("headline"):
            doc.add_paragraph(str(header["headline"]))
        if header.get("contact_line"):
            doc.add_paragraph(str(header["contact_line"]))

        if sections.get("summary"):
            doc.add_heading(labels["summary"], level=1)
            doc.add_paragraph(str(sections["summary"]))

        if sections.get("experience_blocks"):
            doc.add_heading(labels["experience"], level=1)
        for exp in sections.get("experience_blocks") or []:
            doc.add_heading(f"{exp.get('company')} — {exp.get('title')}", level=2)
            meta = " · ".join(
                str(exp[key]) for key in ("date_range", "location") if exp.get(key)
            )
            if meta:
                doc.add_paragraph(meta)
            for bullet in exp.get("bullets", []):
                doc.add_paragraph(str(bullet), style="List Bullet")

        if sections.get("project_blocks"):
            doc.add_heading(labels["projects"], level=1)
        for project in sections.get("project_blocks") or []:
            doc.add_heading(str(project.get("name", "")), level=2)
            for key in ("date_range", "description"):
                if project.get(key):
                    doc.add_paragraph(str(project[key]))
            for bullet in project.get("bullets") or []:
                doc.add_paragraph(str(bullet), style="List Bullet")

        if sections.get("education_blocks"):
            doc.add_heading(labels["education"], level=1)
        for education in sections.get("education_blocks") or []:
            doc.add_heading(str(education.get("school", "")), level=2)
            meta = " · ".join(
                str(education[key])
                for key in ("degree_line", "date_range")
                if education.get(key)
            )
            if meta:
                doc.add_paragraph(meta)

        skills = sections.get("skills_block") or {}
        if skills.get("content"):
            doc.add_heading(str(skills.get("title") or labels["skills"]), level=1)
            doc.add_paragraph(str(skills["content"]))
